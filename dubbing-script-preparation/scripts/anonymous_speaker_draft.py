#!/usr/bin/env python3
"""Build an anonymous-speaker script draft from diarization turns and a trusted SRT."""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import wave
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable


WINDOWS_FFMPEG = Path(r"E:\AI_Models\RVC\RVC20240604Nvidia50x0\ffmpeg.exe")


def default_ffmpeg() -> str:
    configured = os.environ.get("FFMPEG_BINARY", "").strip()
    if configured:
        return configured
    if os.name == "nt" and WINDOWS_FFMPEG.is_file():
        return str(WINDOWS_FFMPEG)
    return "ffmpeg"


TIME_RE = re.compile(
    r"^(?P<h>\d{2}):(?P<m>\d{2}):(?P<s>\d{2})[,.](?P<ms>\d{3})$"
)


class DraftError(RuntimeError):
    pass


def parse_timecode(value: str) -> float:
    match = TIME_RE.match(value.strip())
    if not match:
        raise DraftError(f"Invalid timecode: {value!r}")
    parts = {key: int(number) for key, number in match.groupdict().items()}
    return parts["h"] * 3600 + parts["m"] * 60 + parts["s"] + parts["ms"] / 1000


def format_timecode(seconds: float) -> str:
    milliseconds = max(0, round(seconds * 1000))
    hours, milliseconds = divmod(milliseconds, 3_600_000)
    minutes, milliseconds = divmod(milliseconds, 60_000)
    whole_seconds, milliseconds = divmod(milliseconds, 1000)
    return f"{hours:02d}:{minutes:02d}:{whole_seconds:02d},{milliseconds:03d}"


def parse_srt(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8-sig").replace("\r\n", "\n")
    blocks = [block for block in re.split(r"\n{2,}", text.strip()) if block.strip()]
    rows: list[dict[str, Any]] = []
    for block in blocks:
        lines = block.splitlines()
        if len(lines) < 3:
            raise DraftError(f"Malformed SRT block: {block[:80]!r}")
        try:
            source_index = int(lines[0].strip())
        except ValueError as exc:
            raise DraftError(f"Invalid SRT index: {lines[0]!r}") from exc
        if "-->" not in lines[1]:
            raise DraftError(f"Missing SRT timing arrow at index {source_index}")
        start_text, end_text = [part.strip() for part in lines[1].split("-->", 1)]
        rows.append({
            "source_index": source_index,
            "start": parse_timecode(start_text),
            "end": parse_timecode(end_text),
            "start_timecode": start_text.replace(".", ","),
            "end_timecode": end_text.replace(".", ","),
            "text": "\n".join(lines[2:]),
        })
    return rows


def read_turns_tsv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle, delimiter="\t"))
    turns: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        speaker = str(row.get("speaker") or row.get("speaker_label") or "").strip()
        if not speaker:
            raise DraftError(f"Turn row {row_number} has no speaker")
        try:
            start = float(row.get("start_seconds") or row.get("start") or "")
            end = float(row.get("end_seconds") or row.get("end") or "")
        except ValueError as exc:
            raise DraftError(f"Turn row {row_number} has invalid seconds") from exc
        if end <= start:
            raise DraftError(f"Turn row {row_number} has end <= start")
        turns.append({"start": start, "end": end, "speaker": speaker})
    return sorted(turns, key=lambda item: (item["start"], item["end"], item["speaker"]))


def read_rttm(path: Path) -> list[dict[str, Any]]:
    turns: list[dict[str, Any]] = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip() or line.lstrip().startswith("#"):
            continue
        fields = line.split()
        if len(fields) < 8 or fields[0] != "SPEAKER":
            raise DraftError(f"Unsupported RTTM row {line_number}")
        start = float(fields[3])
        duration = float(fields[4])
        if duration <= 0:
            continue
        turns.append({"start": start, "end": start + duration, "speaker": fields[7]})
    return sorted(turns, key=lambda item: (item["start"], item["end"], item["speaker"]))


def normalize_speakers(turns: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, str]]:
    first_seen: dict[str, float] = {}
    for turn in turns:
        first_seen.setdefault(str(turn["speaker"]), float(turn["start"]))
    ordered = sorted(first_seen, key=lambda speaker: (first_seen[speaker], speaker))
    mapping = {speaker: f"Speaker{index:02d}" for index, speaker in enumerate(ordered, start=1)}
    normalized = [
        {**turn, "source_speaker": turn["speaker"], "speaker": mapping[str(turn["speaker"])]}
        for turn in turns
    ]
    return normalized, mapping


def overlap(start_a: float, end_a: float, start_b: float, end_b: float) -> float:
    return max(0.0, min(end_a, end_b) - max(start_a, start_b))


def map_cues(
    episode: str,
    cues: list[dict[str, Any]],
    turns: list[dict[str, Any]],
    dominance_threshold: float,
) -> list[dict[str, Any]]:
    mapped: list[dict[str, Any]] = []
    for cue in cues:
        by_speaker: dict[str, float] = defaultdict(float)
        for turn in turns:
            amount = overlap(cue["start"], cue["end"], turn["start"], turn["end"])
            if amount:
                by_speaker[str(turn["speaker"])] += amount
        ranked = sorted(by_speaker.items(), key=lambda item: (-item[1], item[0]))
        total = sum(value for _, value in ranked)
        candidates = [speaker for speaker, _ in ranked]
        dominant_ratio = ranked[0][1] / total if ranked and total else 0.0
        if not ranked:
            label = "UNRESOLVED"
            status = "no_speech_turn_overlap"
        elif len(ranked) == 1 or dominant_ratio >= dominance_threshold:
            label = ranked[0][0]
            status = "single_speaker" if len(ranked) == 1 else "dominant_speaker_review"
        else:
            label = "MULTI_SPEAKER_REVIEW"
            status = "multiple_speakers_inside_subtitle"
        mapped.append({
            "episode": episode,
            "source_index": cue["source_index"],
            "start_timecode": cue["start_timecode"],
            "end_timecode": cue["end_timecode"],
            "anonymous_speaker": label,
            "speaker_candidates": "|".join(candidates),
            "dominant_overlap_ratio": f"{dominant_ratio:.4f}",
            "speaker_change_inside_cue": "true" if len(ranked) > 1 else "false",
            "diarization_status": status,
            "text": cue["text"],
        })
    return mapped


def write_tsv(path: Path, rows: Iterable[dict[str, Any]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t")
        writer.writeheader()
        writer.writerows({field: row.get(field, "") for field in fields} for row in rows)


def write_workbook(path: Path, rows: list[dict[str, Any]], title: str = "匿名Speaker草稿") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise DraftError("openpyxl is required to create the anonymous speaker workbook") from exc
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    headers = ["話数", "Speaker", "タイムコード", "台詞", "候補", "要確認"]
    sheet.append(headers)
    for row in rows:
        review = "要確認" if row["diarization_status"] != "single_speaker" else ""
        sheet.append([
            f"第{str(row['episode']).zfill(4)}話",
            row["anonymous_speaker"],
            f"{row['start_timecode']} --> {row['end_timecode']}",
            row["text"],
            row["speaker_candidates"],
            review,
        ])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="385723")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = {"A": 12, "B": 23, "C": 28, "D": 58, "E": 28, "F": 12}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_outputs(
    episode: str,
    srt: Path,
    turns: list[dict[str, Any]],
    out_dir: Path,
    dominance_threshold: float,
) -> dict[str, Any]:
    cues = parse_srt(srt)
    normalized_turns, source_mapping = normalize_speakers(turns)
    mapped = map_cues(episode, cues, normalized_turns, dominance_threshold)
    turn_fields = ["episode", "start_seconds", "end_seconds", "start_timecode", "end_timecode", "speaker", "source_speaker"]
    turn_rows = [{
        "episode": episode,
        "start_seconds": f"{turn['start']:.3f}",
        "end_seconds": f"{turn['end']:.3f}",
        "start_timecode": format_timecode(turn["start"]),
        "end_timecode": format_timecode(turn["end"]),
        "speaker": turn["speaker"],
        "source_speaker": turn["source_speaker"],
    } for turn in normalized_turns]
    cue_fields = [
        "episode", "source_index", "start_timecode", "end_timecode",
        "anonymous_speaker", "speaker_candidates", "dominant_overlap_ratio",
        "speaker_change_inside_cue", "diarization_status", "text",
    ]
    out_dir.mkdir(parents=True, exist_ok=True)
    turns_path = out_dir / "speaker_turns.tsv"
    map_path = out_dir / "subtitle_speaker_map.tsv"
    script_path = out_dir / "anonymous_script.tsv"
    workbook_path = out_dir / "anonymous_speaker_script.xlsx"
    report_path = out_dir / "diarization_report.json"
    write_tsv(turns_path, turn_rows, turn_fields)
    write_tsv(map_path, mapped, cue_fields)
    write_tsv(script_path, mapped, cue_fields)
    write_workbook(workbook_path, mapped)
    review_rows = [row for row in mapped if row["diarization_status"] != "single_speaker"]
    report = {
        "schema_version": 1,
        "episode": episode,
        "source_srt": str(srt.resolve()),
        "subtitle_rows": len(cues),
        "speaker_turns": len(normalized_turns),
        "anonymous_speakers": sorted(source_mapping.values()),
        "source_speaker_mapping": source_mapping,
        "review_rows": len(review_rows),
        "ready": bool(cues) and bool(normalized_turns),
        "identity_claim": False,
        "speaker_scope": "episode_local",
        "rule": "Speaker labels describe acoustic clusters only and are not character identities.",
        "outputs": {
            "speaker_turns": str(turns_path.resolve()),
            "subtitle_speaker_map": str(map_path.resolve()),
            "anonymous_script_tsv": str(script_path.resolve()),
            "anonymous_script_xlsx": str(workbook_path.resolve()),
        },
    }
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


def iter_annotation(annotation: Any) -> Iterable[tuple[float, float, str]]:
    diarization = getattr(annotation, "speaker_diarization", annotation)
    if hasattr(diarization, "itertracks"):
        for segment, _, speaker in diarization.itertracks(yield_label=True):
            yield float(segment.start), float(segment.end), str(speaker)
        return
    raise DraftError("The diarization pipeline returned an unsupported annotation object")


def resolve_model_path(raw_path: str, config_path: Path) -> Path:
    """Resolve relative paths and Windows drive paths when the script runs in WSL."""
    value = raw_path.strip()
    if os.name != "nt" and re.fullmatch(r"[A-Za-z]:[\\/].+", value):
        drive = value[0].lower()
        value = f"/mnt/{drive}/{value[3:].replace(chr(92), '/')}"
    candidate = Path(value).expanduser()
    if not candidate.is_absolute():
        candidate = config_path.parent / candidate
    return candidate.resolve()


def resolve_device(device_name: str) -> str:
    import torch
    requested = device_name.strip().lower()
    if requested == "auto":
        return "cuda" if torch.cuda.is_available() else "cpu"
    if requested.startswith("cuda") and not torch.cuda.is_available():
        raise DraftError("CUDA was requested but PyTorch cannot access a GPU")
    return requested


def load_wespeaker_pipeline(path: Path, config_path: Path, device_name: str = "auto") -> dict[str, Any]:
    """Load the public WeSpeaker ONNX + Silero VAD local diarization stack."""
    try:
        import onnxruntime as ort
        from silero_vad import load_silero_vad
        payload = json.loads(config_path.read_text(encoding="utf-8"))
        if payload.get("backend") != "wespeaker_onnx_spectral":
            raise DraftError("Unsupported anonymous diarization backend")
        if payload.get("clusterer") != "agglomerative_cosine":
            raise DraftError("Anonymous diarization clusterer must be agglomerative_cosine")
        from sklearn.cluster import AgglomerativeClustering  # noqa: F401
        device = resolve_device(device_name)
        model_dir = resolve_model_path(str(payload.get("model_dir") or ""), config_path)
        onnx_model = (model_dir / str(payload.get("onnx_model") or "")).resolve()
        if not model_dir.is_dir() or not onnx_model.is_relative_to(model_dir) or not onnx_model.is_file():
            raise DraftError("WeSpeaker ONNX model is missing or outside model_dir")
        available = set(ort.get_available_providers())
        providers = (["CUDAExecutionProvider", "CPUExecutionProvider"]
                     if device.startswith("cuda") and "CUDAExecutionProvider" in available
                     else ["CPUExecutionProvider"])
        if device.startswith("cuda") and providers[0] != "CUDAExecutionProvider":
            raise DraftError("CUDA was requested but ONNX Runtime has no CUDAExecutionProvider")
        options = ort.SessionOptions()
        options.inter_op_num_threads = 1
        options.intra_op_num_threads = 1
        session = ort.InferenceSession(str(onnx_model), sess_options=options, providers=providers)
        if [item.name for item in session.get_inputs()] != ["feats"] or [item.name for item in session.get_outputs()] != ["embs"]:
            raise DraftError("Unexpected WeSpeaker ONNX input/output contract")
        return {
            "backend": "wespeaker_onnx_spectral",
            "root": str(path),
            "config": payload,
            "session": session,
            "vad": load_silero_vad(onnx=True),
            "device": device,
            "providers": session.get_providers(),
        }
    except DraftError:
        raise
    except Exception as exc:
        raise DraftError(f"Local WeSpeaker diarization pipeline could not load: {exc}") from exc


def load_local_pipeline(path: Path, device_name: str = "auto") -> Any:
    """Resolve a fully local pipeline and disable Hub downloads before loading it."""
    path = path.resolve()
    backend_config = path / "backend_config.json" if path.is_dir() else path
    if backend_config.is_file() and backend_config.name == "backend_config.json":
        return load_wespeaker_pipeline(path, backend_config, device_name)
    config = next((item for item in (path / "config.yaml", path / "config.yml") if item.is_file()), path) if path.is_dir() else path
    if not config.is_file():
        raise DraftError(f"Local diarization config not found: {config}")
    os.environ["HF_HUB_OFFLINE"] = "1"
    os.environ["TRANSFORMERS_OFFLINE"] = "1"
    try:
        import yaml
        import numpy
        import torch
        from pyannote.audio import Pipeline
        payload = yaml.safe_load(config.read_text(encoding="utf-8"))
        parameters = (payload.get("pipeline") or {}).get("params") or {}
        for name in ("segmentation", "embedding"):
            value = parameters.get(name)
            local = Path(str(value or ""))
            if not value or not local.is_absolute() or not local.exists():
                raise DraftError(f"Pipeline {name} must reference an existing absolute local model path; remote model IDs are disabled")
        pipeline = Pipeline.from_pretrained(str(config))
        if pipeline is None:
            raise DraftError("Local diarization pipeline did not load")
        return pipeline
    except DraftError:
        raise
    except Exception as exc:
        raise DraftError(f"Local diarization pipeline could not load: {exc}") from exc


def spectral_cluster(embeddings: Any, num_speakers: int | None, min_speakers: int, max_speakers: int) -> list[int]:
    """WeSpeaker spectral clustering without its kaldiio training dependency."""
    import numpy as np
    import scipy.linalg
    from sklearn.cluster import KMeans

    matrix = np.asarray(embeddings, dtype=np.float32)
    if len(matrix) <= 2:
        return [0] * len(matrix)
    norms = np.linalg.norm(matrix, axis=1, keepdims=True)
    matrix = matrix / np.maximum(norms, 1e-8)
    similarity = 0.5 * (1.0 + matrix @ matrix.T)
    count = similarity.shape[0]
    retain_from = max(count - 10, 2) if count < 1000 else int(0.99 * count)
    for row in range(count):
        indexes = np.argsort(similarity[row])
        similarity[row, indexes[:retain_from]] = 0.0
        similarity[row, indexes[retain_from:]] = 1.0
    similarity = 0.5 * (similarity + similarity.T)
    similarity[np.diag_indices(count)] = 0.0
    laplacian = np.diag(np.sum(np.abs(similarity), axis=1)) - similarity
    eigenvalues, eigenvectors = scipy.linalg.eigh(laplacian)
    upper = max(1, min(max_speakers, count - 1))
    if num_speakers is None:
        gaps = np.diff(eigenvalues[:upper + 1])
        clusters = int(np.argmax(gaps)) + 1 if len(gaps) else 1
    else:
        clusters = int(num_speakers)
    clusters = max(min_speakers, min(clusters, upper))
    if clusters == 1:
        return [0] * count
    labels = KMeans(n_clusters=clusters, random_state=0, n_init=10).fit_predict(eigenvectors[:, :clusters])
    return [int(value) for value in labels]


def agglomerative_cosine_cluster(embeddings: Any, similarity_threshold: float) -> list[int]:
    """Cluster unknown speaker counts with a reviewable cosine threshold."""
    import numpy as np
    from sklearn.cluster import AgglomerativeClustering

    matrix = np.asarray(embeddings, dtype=np.float32)
    if len(matrix) <= 2:
        return [0] * len(matrix)
    matrix = matrix / np.maximum(np.linalg.norm(matrix, axis=1, keepdims=True), 1e-8)
    labels = AgglomerativeClustering(
        n_clusters=None,
        metric="cosine",
        linkage="average",
        distance_threshold=1.0 - similarity_threshold,
    ).fit_predict(matrix)
    return [int(value) for value in labels]


def merge_labeled_segments(segments: list[tuple[float, float, int]]) -> list[dict[str, Any]]:
    if not segments:
        return []
    begin, end, label = segments[0]
    merged: list[dict[str, Any]] = []
    for next_begin, next_end, next_label in segments[1:]:
        if next_begin <= end and next_label == label:
            end = max(end, next_end)
        elif next_begin > end:
            merged.append({"start": begin, "end": end, "speaker": f"cluster_{label}"})
            begin, end, label = next_begin, next_end, next_label
        else:
            pivot = (next_begin + end) / 2.0
            merged.append({"start": begin, "end": pivot, "speaker": f"cluster_{label}"})
            begin, end, label = pivot, next_end, next_label
    merged.append({"start": begin, "end": end, "speaker": f"cluster_{label}"})
    return merged


def run_wespeaker_diarization(pipeline: dict[str, Any], waveform: Any, sample_rate: int) -> list[dict[str, Any]]:
    import numpy as np
    import torch
    import torchaudio.compliance.kaldi as kaldi
    from silero_vad import get_speech_timestamps

    if sample_rate != 16000:
        raise DraftError("WeSpeaker anonymous diarization expects 16000 Hz audio")
    config = pipeline["config"]
    min_duration_ms = int(config.get("min_speech_duration_ms", 255))
    speech = get_speech_timestamps(
        waveform.squeeze(0),
        pipeline["vad"],
        sampling_rate=sample_rate,
        min_speech_duration_ms=min_duration_ms,
        return_seconds=True,
    )
    window_frames = int(float(config.get("window_seconds", 1.5)) * 100)
    period_frames = int(float(config.get("period_seconds", 0.75)) * 100)
    subsegs: list[tuple[float, float]] = []
    feature_windows: list[Any] = []
    for segment in speech:
        begin = float(segment["start"])
        end = float(segment["end"])
        if (end - begin) * 1000 < min_duration_ms:
            continue
        chunk = waveform[:, int(begin * sample_rate):int(end * sample_rate)].to(torch.float32)
        fbank = kaldi.fbank(
            chunk,
            num_mel_bins=80,
            frame_length=25,
            frame_shift=10,
            sample_frequency=sample_rate,
            window_type="hamming",
        ).numpy()
        nominal_frames = max(1, int((end - begin) * 100))
        starts = [0] if nominal_frames <= window_frames else list(range(0, nominal_frames - window_frames + period_frames, period_frames))
        for start_frame in starts:
            end_frame = min(start_frame + window_frames, nominal_frames)
            window = np.resize(fbank[start_frame:end_frame], (window_frames, 80)).astype("float32", copy=False)
            window -= np.mean(window, axis=0, keepdims=True)
            feature_windows.append(window)
            subsegs.append((begin + start_frame / 100.0, begin + end_frame / 100.0))
    if not feature_windows:
        return []
    features = np.stack(feature_windows)
    embeddings = pipeline["session"].run(["embs"], {"feats": features})[0]
    configured = config.get("num_speakers")
    if configured not in (None, ""):
        labels = spectral_cluster(
            embeddings,
            int(configured),
            int(config.get("min_speakers", 1)),
            int(config.get("max_speakers", 20)),
        )
    elif config.get("clusterer") == "agglomerative_cosine":
        threshold = float(config.get("speaker_similarity_threshold", 0.70))
        if not 0.0 < threshold < 1.0:
            raise DraftError("speaker_similarity_threshold must be between 0 and 1")
        labels = agglomerative_cosine_cluster(embeddings, threshold)
    else:
        labels = spectral_cluster(
            embeddings,
            None,
            int(config.get("min_speakers", 1)),
            int(config.get("max_speakers", 20)),
        )
    return merge_labeled_segments([(begin, end, label) for (begin, end), label in zip(subsegs, labels)])


def run_diarization(args: argparse.Namespace) -> list[dict[str, Any]]:
    media_path = Path(args.media).resolve()
    if not media_path.is_file():
        raise DraftError(f"Speech-dominant media not found: {media_path}")
    import numpy as np
    import torch
    device = resolve_device(args.device)
    pipeline = load_local_pipeline(Path(args.pipeline), device)
    if not isinstance(pipeline, dict):
        pipeline.to(torch.device(device))
    ffmpeg = Path(args.ffmpeg).resolve()
    if not ffmpeg.is_file():
        raise DraftError(f"ffmpeg not found: {ffmpeg}")
    with tempfile.TemporaryDirectory(prefix="anonymous_diarization_") as temp_name:
        decoded = Path(temp_name) / "speech.wav"
        completed = subprocess.run(
            [
                str(ffmpeg), "-hide_banner", "-loglevel", "error", "-y",
                "-i", str(media_path), "-vn", "-ac", "1", "-ar", "16000",
                "-c:a", "pcm_s16le", str(decoded),
            ],
            check=False,
            capture_output=True,
            text=True,
        )
        if completed.returncode != 0 or not decoded.is_file():
            raise DraftError(f"ffmpeg could not decode diarization media: {completed.stderr.strip()}")
        with wave.open(str(decoded), "rb") as handle:
            sample_rate = handle.getframerate()
            channels = handle.getnchannels()
            width = handle.getsampwidth()
            raw = handle.readframes(handle.getnframes())
        if channels != 1 or width != 2:
            raise DraftError("Decoded diarization audio is not mono 16-bit PCM")
        samples = np.frombuffer(raw, dtype="<i2").astype("float32") / 32768.0
        waveform = torch.from_numpy(samples).unsqueeze(0)
        if isinstance(pipeline, dict):
            return run_wespeaker_diarization(pipeline, waveform, sample_rate)
        annotation = pipeline({"waveform": waveform, "sample_rate": sample_rate})
    return [
        {"start": start, "end": end, "speaker": speaker}
        for start, end, speaker in iter_annotation(annotation)
        if end > start
    ]


def cmd_build(args: argparse.Namespace) -> int:
    if bool(args.turns_tsv) == bool(args.rttm):
        raise DraftError("Provide exactly one of --turns-tsv or --rttm")
    turns = read_turns_tsv(Path(args.turns_tsv)) if args.turns_tsv else read_rttm(Path(args.rttm))
    report = build_outputs(args.episode, Path(args.srt), turns, Path(args.out_dir), args.dominance_threshold)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ready"] else 2


def cmd_diarize(args: argparse.Namespace) -> int:
    turns = run_diarization(args)
    report = build_outputs(args.episode, Path(args.srt), turns, Path(args.out_dir), args.dominance_threshold)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ready"] else 2


def cmd_preflight(args: argparse.Namespace) -> int:
    checks: list[dict[str, Any]] = []
    pipeline = Path(args.pipeline).resolve()
    try:
        loaded_pipeline = load_local_pipeline(pipeline, args.device)
        pipeline_runtime = {
            "device": loaded_pipeline.get("device", resolve_device(args.device)) if isinstance(loaded_pipeline, dict) else resolve_device(args.device),
            "providers": loaded_pipeline.get("providers", []) if isinstance(loaded_pipeline, dict) else [],
        }
        del loaded_pipeline
        checks.append({"component": "anonymous_diarization_pipeline", "path": str(pipeline), "ready": True, **pipeline_runtime})
    except DraftError as exc:
        checks.append({"component": "anonymous_diarization_pipeline", "path": str(pipeline), "ready": False, "error": str(exc)})
    required_names = {"campplus", "ecapa512", "resnet34"}
    provided: dict[str, Path] = {}
    for value in args.voice_model:
        if "=" not in value:
            raise DraftError("--voice-model must use name=/absolute/path")
        name, raw_path = value.split("=", 1)
        name = name.strip().casefold()
        if name not in required_names or name in provided:
            raise DraftError(f"Unknown or duplicate identity model: {name}")
        provided[name] = Path(raw_path).resolve()
    if len(set(provided.values())) != len(provided):
        raise DraftError("Each identity encoder must use a distinct model directory")
    main_skill = Path(args.main_skill_dir).resolve() if args.main_skill_dir else Path(__file__).resolve().parents[2] / "dubbing-script-automation"
    voice_tool = main_skill / "scripts" / "voice_evidence.py"
    for name in sorted(required_names):
        path = provided.get(name)
        item = {"component": f"voice_identity_{name}", "path": str(path) if path else "", "ready": False}
        if path and path.is_dir() and any(path.iterdir()) and voice_tool.is_file():
            command = [sys.executable, str(voice_tool), "preflight", "--model-kind", name, "--model", str(path), "--ffmpeg", args.ffmpeg, "--device", args.device]
            completed = subprocess.run(command, capture_output=True, text=True, check=False, env={**os.environ, "HF_HUB_OFFLINE": "1", "TRANSFORMERS_OFFLINE": "1"})
            try:
                report = json.loads(completed.stdout)
            except json.JSONDecodeError:
                report = {}
            item["ready"] = completed.returncode == 0 and (report.get("runtime_ready") is True or report.get("ready") is True)
            if not item["ready"]:
                item["error"] = completed.stderr.strip() or completed.stdout.strip() or "identity model preflight failed"
        else:
            item["error"] = "identity model directory is missing/empty or preflight tool is missing"
        checks.append(item)
    ffmpeg = shutil.which(args.ffmpeg) if not Path(args.ffmpeg).is_absolute() else args.ffmpeg
    checks.append({"component": "ffmpeg", "path": str(ffmpeg or ""), "ready": bool(ffmpeg and Path(ffmpeg).is_file() and os.access(ffmpeg, os.X_OK))})
    ready = all(item["ready"] for item in checks)
    payload = {"ready": ready, "wavlm_used": False, "required_voice_identity_models": sorted(required_names), "checks": checks}
    if args.report:
        Path(args.report).write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0 if ready else 2


def cmd_merge(args: argparse.Namespace) -> int:
    root = Path(args.root).resolve()
    rows: list[dict[str, Any]] = []
    selected = {value.strip() for value in (getattr(args, "episodes", "") or "").split(",") if value.strip()}
    if any(not re.fullmatch(r"[0-9]{4}", value) for value in selected):
        raise DraftError("--episodes must contain four-digit episode IDs")
    paths = [root / episode / "anonymous_script.tsv" for episode in sorted(selected)] if selected else sorted(root.glob("*/anonymous_script.tsv"))
    for path in paths:
        if not path.is_file():
            raise DraftError(f"Missing requested anonymous episode draft: {path}")
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows.extend(csv.DictReader(handle, delimiter="\t"))
    if not rows:
        raise DraftError(f"No episode anonymous_script.tsv files found under: {root}")
    rows.sort(key=lambda row: (str(row.get("episode", "")), int(row.get("source_index", 0))))
    fields = [
        "episode", "source_index", "start_timecode", "end_timecode",
        "anonymous_speaker", "speaker_candidates", "dominant_overlap_ratio",
        "speaker_change_inside_cue", "diarization_status", "text",
    ]
    out_tsv = Path(args.out_tsv).resolve()
    out_xlsx = Path(args.out_xlsx).resolve()
    write_tsv(out_tsv, rows, fields)
    write_workbook(out_xlsx, rows, title="全話匿名Speaker草稿")
    payload = {
        "episodes": sorted({str(row["episode"]) for row in rows}),
        "rows": len(rows),
        "out_tsv": str(out_tsv),
        "out_xlsx": str(out_xlsx),
        "identity_claim": False,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Create a no-character-name anonymous speaker script draft")
    subparsers = parser.add_subparsers(dest="command", required=True)
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--episode", required=True)
    common.add_argument("--srt", required=True)
    common.add_argument("--out-dir", required=True)
    common.add_argument("--dominance-threshold", type=float, default=0.70)
    build = subparsers.add_parser("build", parents=[common], help="Build from reviewed turns TSV or RTTM")
    build.add_argument("--turns-tsv")
    build.add_argument("--rttm")
    build.set_defaults(func=cmd_build)
    diarize = subparsers.add_parser("diarize", parents=[common], help="Run a local WeSpeaker or pyannote diarization pipeline and build the draft")
    diarize.add_argument("--media", required=True)
    diarize.add_argument("--pipeline", required=True)
    diarize.add_argument("--device", default="auto")
    diarize.add_argument("--ffmpeg", default=default_ffmpeg())
    diarize.set_defaults(func=cmd_diarize)
    preflight = subparsers.add_parser("preflight", help="Check the anonymous-segmentation and three-model identity stack")
    preflight.add_argument("--pipeline", required=True)
    preflight.add_argument("--main-skill-dir")
    preflight.add_argument("--voice-model", action="append", default=[])
    preflight.add_argument("--ffmpeg", default=default_ffmpeg())
    preflight.add_argument("--device", default="auto")
    preflight.add_argument("--report")
    preflight.set_defaults(func=cmd_preflight)
    merge = subparsers.add_parser("merge", help="Merge per-episode anonymous drafts into one internal whole-series draft")
    merge.add_argument("--root", required=True)
    merge.add_argument("--episodes", help="Comma-separated exact episode scope; excludes stale drafts from older builds")
    merge.add_argument("--out-tsv", required=True)
    merge.add_argument("--out-xlsx", required=True)
    merge.set_defaults(func=cmd_merge)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    try:
        if not 0.5 <= getattr(args, "dominance_threshold", 0.70) <= 1.0:
            raise DraftError("--dominance-threshold must be between 0.5 and 1.0")
        return args.func(args)
    except (OSError, ValueError, DraftError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
