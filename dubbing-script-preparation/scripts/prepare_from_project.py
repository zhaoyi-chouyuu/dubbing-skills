#!/usr/bin/env python3
"""Discover a dubbing project's source material and build its preparation package."""

from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from typing import Any


VIDEO_SUFFIXES = {".mp4", ".mov", ".mkv", ".m4v"}
SUBTITLE_SUFFIXES = {".srt", ".ass", ".ssa", ".vtt"}
WORKBOOK_SUFFIXES = {".xlsx", ".xls", ".csv", ".tsv"}
AUDIO_SUFFIXES = {".wav", ".mp3", ".m4a", ".flac", ".aac", ".ogg"}
DIRECTORY_ROLE_HINTS = {
    "feature": {"feature"},
    "finish": {"finish", "delivery", "deliverable"},
    "video": {"video", "movie", "footage"},
    "subtitle": {"subtitle", "subtitles", "caption", "captions", "srt"},
    "script": {"script", "scenario"},
    "casting": {"casting", "cast", "character", "characters", "role", "roles"},
    "sound": {"sound", "audio"},
    "vocal": {"vocal", "vocals", "voice", "voices"},
    "fix": {"fix", "fixed", "revision", "revised", "corrected"},
}
IGNORED_DIRS = {
    "01_preparation", "02_script_work", "03_final_delivery",
    "前期准备包", "脚本作业", "正式脚本", "05_全剧交付",
}
VISUAL_VIDEO_PRIORITY = ("01_original", "02_dry_video", "03_BGM_video", "03_subtitiled")
SPEECH_DOMINANT_VIDEO_PRIORITY = ("02_dry_video",)
ROLE_HEADERS = {"ID", "性别", "年龄", "身份", "人物描述", "声线描述"}
ROLE_NAME_HEADERS = {"角色名", "角色", "姓名", "名称", "人物", "役名", "名前", "name", "role", "character"}
GENERATED_CAST_WORKBOOK_MARKERS = (
    "主要人物角色表",
    "主要登場人物一覧",
    "主役登場人物",
)


def directory_tokens(path: Path) -> set[str]:
    """Classify numbered project folders by their English semantic labels."""
    return set(re.findall(r"[a-z]+", path.name.casefold()))


def directory_file_signals(path: Path, max_depth: int = 2) -> dict[str, int]:
    signals = {"video": 0, "subtitle": 0, "workbook": 0, "audio": 0}
    pending = [(path, 0)]
    while pending:
        current, depth = pending.pop()
        try:
            children = list(current.iterdir())
        except OSError:
            continue
        for item in children:
            if item.name.startswith(".") or item.name in IGNORED_DIRS:
                continue
            if item.is_dir():
                if depth < max_depth:
                    pending.append((item, depth + 1))
                continue
            if not item.is_file():
                continue
            suffix = item.suffix.casefold()
            if suffix in VIDEO_SUFFIXES:
                signals["video"] += 1
            if suffix in SUBTITLE_SUFFIXES:
                signals["subtitle"] += 1
            if suffix in WORKBOOK_SUFFIXES:
                signals["workbook"] += 1
            if suffix in AUDIO_SUFFIXES:
                signals["audio"] += 1
    return signals


def infer_directory_roles(path: Path) -> tuple[list[str], dict[str, int]]:
    """Infer folder purpose from names and bounded content; hints are not a whitelist."""
    tokens = directory_tokens(path)
    roles = {
        role for role, hints in DIRECTORY_ROLE_HINTS.items()
        if tokens & hints
    }
    signals = directory_file_signals(path)
    if signals["video"]:
        roles.add("video")
    if signals["subtitle"]:
        roles.add("subtitle")
    if signals["audio"] and not ({"sound", "vocal"} & roles):
        roles.add("sound")
    if signals["workbook"] and any(
        item.stem == "画面字" for item in path.glob("*.xlsx")
    ):
        roles.add("script")
    return sorted(roles), signals


def directory_inventory(parent: Path) -> list[dict[str, Any]]:
    inventory: list[dict[str, Any]] = []
    for path in sorted(item for item in parent.iterdir() if item.is_dir() and not item.name.startswith(".")):
        roles, signals = infer_directory_roles(path)
        inventory.append({
            "path": str(path.resolve()),
            "name_tokens": sorted(directory_tokens(path)),
            "inferred_roles": roles,
            "file_signals": signals,
            "recognized": bool(roles),
        })
    return inventory


def child_directory(parent: Path, role: str, *, required: bool = True) -> Path | None:
    exact = parent / role
    if exact.is_dir():
        return exact.resolve()
    children = [path for path in sorted(parent.iterdir()) if path.is_dir() and not path.name.startswith(".")]
    named_matches = [
        path.resolve() for path in children
        if directory_tokens(path) & DIRECTORY_ROLE_HINTS.get(role, {role.casefold()})
    ]
    if len(named_matches) == 1:
        return named_matches[0]
    matches = [
        path.resolve() for path in children
        if role in infer_directory_roles(path)[0]
        and "finish" not in infer_directory_roles(path)[0]
    ]
    if len(matches) == 1:
        return matches[0]
    if required:
        raise ValueError(
            f"Could not uniquely infer the {role} directory under {parent}; found {len(matches)} candidates"
        )
    return None


def is_ignored(path: Path) -> bool:
    return path.name.startswith("._") or any(part.startswith(".") or part in IGNORED_DIRS for part in path.parts)


def is_generated_cast_workbook(path: Path) -> bool:
    """Return true for production-facing cast summaries, not source role workbooks."""
    return (
        path.suffix.casefold() == ".xlsx"
        and any(marker in path.stem for marker in GENERATED_CAST_WORKBOOK_MARKERS)
    )


def episode_id(path: Path) -> str | None:
    matches = re.findall(r"(?:^|[_-])(\d{4})(?=$|[_-])", path.stem)
    return matches[-1] if matches else None


def locate_feature_root(project_root: Path) -> Path:
    project_root = project_root.resolve()
    def has_source_structure(path: Path) -> bool:
        try:
            return all(child_directory(path, role, required=False) for role in ("video", "subtitle", "script"))
        except OSError:
            return False

    try:
        if has_source_structure(project_root):
            return project_root
    except OSError:
        pass
    candidates = [
        path for path in sorted(project_root.iterdir())
        if path.is_dir()
        and "finish" not in infer_directory_roles(path)[0]
        and has_source_structure(path)
    ]
    named = [path for path in candidates if "feature" in infer_directory_roles(path)[0]]
    if len(named) == 1:
        return named[0].resolve()
    if len(candidates) != 1:
        raise ValueError(
            f"Could not uniquely infer the feature/source folder from directory semantics and contents; found {len(candidates)} candidates"
        )
    return candidates[0].resolve()


def choose_named_dir(parent: Path, priority: tuple[str, ...], *, allow_direct: bool = True) -> Path | None:
    for name in priority:
        candidate = parent / name
        if candidate.is_dir():
            return candidate
    children = [path for path in sorted(parent.iterdir()) if path.is_dir()]
    for name in priority:
        required_tokens = set(re.findall(r"[a-z]+", name.casefold()))
        matches = [path for path in children if required_tokens.issubset(directory_tokens(path))]
        if len(matches) == 1:
            return matches[0].resolve()
        if len(matches) > 1:
            raise ValueError(
                f"Ambiguous {name} directory role under {parent}: {', '.join(path.name for path in matches)}"
            )
    direct_files = [path for path in parent.iterdir() if path.is_file() and path.suffix.casefold() in VIDEO_SUFFIXES]
    return parent if allow_direct and direct_files else None


def index_media(root: Path | None, suffixes: set[str]) -> tuple[dict[str, Path], list[str]]:
    if root is None:
        return {}, []
    grouped: dict[str, list[Path]] = {}
    skipped: list[str] = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or is_ignored(path) or path.suffix.casefold() not in suffixes:
            continue
        episode = episode_id(path)
        if episode is None:
            skipped.append(str(path.resolve()))
            continue
        grouped.setdefault(episode, []).append(path.resolve())
    duplicates = [f"episode {episode}: duplicate files: {', '.join(map(str, paths))}" for episode, paths in grouped.items() if len(paths) != 1]
    return {episode: paths[0] for episode, paths in grouped.items() if len(paths) == 1}, [*duplicates, *[f"unrecognized episode id: {path}" for path in skipped]]


def workbook_headers(path: Path) -> set[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("openpyxl is required; run this script with the stable-ai Python runtime") from exc
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        if sheet.max_column is None:
            sheet.calculate_dimension(force=True)
        max_column = sheet.max_column or 1
        return {str(sheet.cell(1, column).value or "").strip() for column in range(1, max_column + 1)}
    finally:
        workbook.close()


def normalize_role_workbook(source: Path, destination: Path) -> dict[str, Any]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("openpyxl is required; run this script with the stable-ai Python runtime") from exc
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = [[str(cell.value or "").strip() for cell in row] for row in sheet.iter_rows()]
    finally:
        workbook.close()
    if not rows:
        raise ValueError(f"Role workbook is empty: {source}")
    headers = [header or f"column_{index + 1}" for index, header in enumerate(rows[0])]
    if len(set(headers)) != len(headers):
        raise ValueError("Role workbook has duplicate column headers")
    normalized_headers = [header.casefold() for header in headers]
    name_column = next((index for index, header in enumerate(normalized_headers) if header in {item.casefold() for item in ROLE_NAME_HEADERS}), None)
    id_column = next((index for index, header in enumerate(normalized_headers) if header == "id"), None)
    if name_column is None and id_column is None:
        raise ValueError("Role workbook has neither a role-name column nor an ID column")
    source_headers = [f"source_{header}" if header in {"角色名", "角色来源状态"} else header for header in headers]
    output_headers = ["角色名", "角色来源状态", *source_headers]
    if len(set(output_headers)) != len(output_headers):
        raise ValueError("Role workbook column names collide with normalized source names")
    normalized_rows: list[dict[str, str]] = []
    for row_number, row in enumerate(rows[1:], start=2):
        row += [""] * (len(headers) - len(row))
        supplied_name = row[name_column].strip() if name_column is not None else ""
        source_id = row[id_column].strip() if id_column is not None else str(row_number)
        if not supplied_name and not source_id:
            continue
        role_name = supplied_name or f"角色ID_{source_id}"
        status = "source_named" if supplied_name else "stable_source_id_pending_identity_review"
        item = {"角色名": role_name, "角色来源状态": status}
        item.update({header: row[index].strip() for index, header in enumerate(source_headers)})
        normalized_rows.append(item)
    if not normalized_rows:
        raise ValueError("Role workbook produced no usable role rows")
    with destination.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=output_headers, delimiter="\t")
        writer.writeheader()
        writer.writerows(normalized_rows)
    return {
        "source": str(source),
        "normalized": str(destination),
        "roles": len(normalized_rows),
        "mode": "source_names" if name_column is not None else "stable_source_ids_pending_identity_review",
    }


def project_name(feature_root: Path) -> str:
    raw = feature_root.parent.name if "feature" in directory_tokens(feature_root) else feature_root.name
    return re.sub(r"[（(][^）)]*[）)]$", "", raw).strip()


def discover_workbooks(feature_root: Path, script_root: Path, explicit_roles: str | None, explicit_screen: str | None) -> tuple[Path, Path, list[dict[str, Any]]]:
    search_roots = [script_root.resolve()]
    for item in directory_inventory(feature_root):
        path = Path(item["path"])
        roles = set(item["inferred_roles"])
        if "finish" in roles or path == script_root.resolve():
            continue
        if "casting" in roles or int(item["file_signals"]["workbook"]) > 0:
            search_roots.append(path)
    candidates: list[Path] = []
    for root in search_roots:
        for pattern in ("*.xlsx", "*/*.xlsx"):
            for path in sorted(root.glob(pattern)):
                resolved = path.resolve()
                if (
                    not is_ignored(path)
                    and not is_generated_cast_workbook(path)
                    and resolved not in candidates
                ):
                    candidates.append(resolved)
    inventory: list[dict[str, Any]] = []
    for path in candidates:
        headers = workbook_headers(path)
        inventory.append({
            "path": str(path),
            "source_directory": str(path.parent),
            "source_directory_roles": infer_directory_roles(path.parent)[0],
            "headers": sorted(headers),
            "role_header_matches": len(ROLE_HEADERS & headers),
        })

    if explicit_roles:
        roles = Path(explicit_roles).resolve()
    else:
        exact_name = project_name(feature_root)
        scored: list[tuple[int, Path]] = []
        for item in inventory:
            path = Path(item["path"])
            score = int(item["role_header_matches"]) * 10
            if path.stem == exact_name:
                score += 100
            if any(marker in path.stem for marker in ("日本語脚本", "総合脚本", "主要人物角色表", "画面字")):
                score -= 100
            scored.append((score, path))
        scored.sort(key=lambda item: (-item[0], item[1].name))
        if not scored or scored[0][0] < 40 or (len(scored) > 1 and scored[0][0] == scored[1][0]):
            raise ValueError("Could not uniquely identify the project role workbook; pass --roles explicitly")
        roles = scored[0][1]

    if explicit_screen:
        screen = Path(explicit_screen).resolve()
    else:
        screen_candidates = [path for path in candidates if path.stem == "画面字"]
        if len(screen_candidates) != 1:
            raise ValueError(f"Expected exactly one 画面字.xlsx in the discovered script/casting material; found {len(screen_candidates)}")
        screen = screen_candidates[0]
    if not roles.is_file() or not screen.is_file():
        raise ValueError("The discovered role workbook or 画面字.xlsx is missing")
    return roles, screen, inventory


def discover(project_root: Path, requested_episodes: set[str] | None = None) -> dict[str, Any]:
    feature_root = locate_feature_root(project_root)
    video_parent = child_directory(feature_root, "video")
    script_root = child_directory(feature_root, "script")
    visual_root = choose_named_dir(video_parent, VISUAL_VIDEO_PRIORITY)
    speech_dominant_root = choose_named_dir(video_parent, SPEECH_DOMINANT_VIDEO_PRIORITY, allow_direct=False)
    subtitle_root = child_directory(feature_root, "subtitle")
    visual, visual_issues = index_media(visual_root, VIDEO_SUFFIXES)
    speech_dominant, speech_dominant_issues = index_media(speech_dominant_root, VIDEO_SUFFIXES)
    subtitles, subtitle_issues = index_media(subtitle_root, {".srt"})
    episodes = sorted(set(visual) | set(subtitles))
    if requested_episodes:
        unknown = sorted(requested_episodes - set(episodes))
        if unknown:
            raise ValueError("Requested episodes not found: " + ", ".join(unknown))
        episodes = [episode for episode in episodes if episode in requested_episodes]
    issues = [*visual_issues, *subtitle_issues]
    rows: list[dict[str, str]] = []
    for episode in episodes:
        if episode not in visual:
            issues.append(f"episode {episode}: video missing")
            continue
        if episode not in subtitles:
            issues.append(f"episode {episode}: srt missing")
            continue
        row = {"episode": episode, "video": str(visual[episode]), "srt": str(subtitles[episode])}
        if speech_dominant_root is not None:
            if episode in speech_dominant:
                row["speech_dominant_audio"] = str(speech_dominant[episode])
            else:
                issues.append(f"episode {episode}: 02_dry_video counterpart missing")
        rows.append(row)
    issues.extend(issue for issue in speech_dominant_issues if issue not in visual_issues)
    return {
        "schema_version": 2,
        "project_root": str(project_root.resolve()),
        "feature_root": str(feature_root),
        "script_root": str(script_root),
        "visual_video_root": str(visual_root) if visual_root else "",
        "speech_dominant_video_root": str(speech_dominant_root) if speech_dominant_root else "",
        "subtitle_root": str(subtitle_root),
        "directory_inventory": directory_inventory(feature_root),
        "directory_discovery_policy": (
            "Known role names are hints, not a whitelist. Inspect unknown directories by name, "
            "hierarchy, file types, and content; require explicit choice only when candidates remain ambiguous."
        ),
        "episode_rows": rows,
        "issues": sorted(set(issues)),
    }


def write_manifest(path: Path, rows: list[dict[str, str]]) -> None:
    fields = ["episode", "video", "srt"]
    if rows and any("speech_dominant_audio" in row for row in rows):
        fields.append("speech_dominant_audio")
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def parse_episode_filter(value: str | None) -> set[str] | None:
    if not value:
        return None
    episodes = {item.strip().zfill(4) for item in value.split(",") if item.strip()}
    return episodes or None


def main() -> int:
    parser = argparse.ArgumentParser(description="Discover project material and build a preparation package inside the feature/script directory")
    parser.add_argument("--project-root", required=True, help="Project root or its inferred source-production folder")
    parser.add_argument("--roles", help="Override automatic project role-workbook discovery")
    parser.add_argument("--screen-text-xlsx", help="Override automatic 画面字.xlsx discovery")
    parser.add_argument("--episodes", help="Optional comma-separated episode IDs, e.g. 0001,0002")
    parser.add_argument("--discover-only", action="store_true", help="Print discovery JSON without creating any files")
    parser.add_argument("--out-dir", help="Override the default <script>/01_preparation name; it must remain a direct child of the detected script directory")
    parser.add_argument("--frames-per-subtitle", type=int, choices=[1, 3], default=3)
    parser.add_argument("--main-skill-dir")
    parser.add_argument(
        "--anonymous-diarization-pipeline",
        help="Local WeSpeaker or pyannote diarization pipeline used to create the mandatory Speaker01/Speaker02 preparation draft",
    )
    parser.add_argument("--anonymous-device", default="auto")
    parser.add_argument("--identity-model", action="append", default=[])
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    try:
        project_root = Path(args.project_root).resolve()
        result = discover(project_root, parse_episode_filter(args.episodes))
        feature_root = Path(result["feature_root"])
        script_root = Path(result["script_root"])
        roles, screen, workbook_inventory = discover_workbooks(
            feature_root, script_root, args.roles, args.screen_text_xlsx,
        )
        result["roles_workbook"] = str(roles)
        result["screen_text_workbook"] = str(screen)
        result["workbook_inventory"] = workbook_inventory
        if not result["episode_rows"]:
            result["issues"].append("no complete video/SRT episode pairs")
        if args.discover_only or result["issues"]:
            print(json.dumps(result, ensure_ascii=False, indent=2))
            return 0 if args.discover_only and not result["issues"] else 2

        package = Path(args.out_dir).resolve() if args.out_dir else script_root / "01_preparation"
        if package.parent.resolve() != script_root.resolve():
            raise ValueError(
                f"Preparation package must be a direct child of the script directory: {script_root}"
            )
        core = Path(__file__).with_name("prepare_package.py")
        with tempfile.TemporaryDirectory(prefix="dubbing-preparation-") as temp_name:
            temp_root = Path(temp_name)
            manifest = temp_root / "episode_manifest.tsv"
            discovery_json = temp_root / "discovered_materials.json"
            normalized_roles = temp_root / "normalized_roles.tsv"
            write_manifest(manifest, result["episode_rows"])
            result["role_catalog"] = normalize_role_workbook(roles, normalized_roles)
            discovery_json.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
            command = [
                sys.executable, str(core),
                "--episode-manifest", str(manifest),
                "--roles", str(normalized_roles),
                "--source-role-workbook", str(roles),
                "--screen-text-xlsx", str(screen),
                "--out-dir", str(package),
                "--feature-root", str(feature_root),
                "--script-root", str(script_root),
                "--frames-per-subtitle", str(args.frames_per_subtitle),
            ]
            if args.main_skill_dir:
                command.extend(["--main-skill-dir", args.main_skill_dir])
            if args.anonymous_diarization_pipeline:
                command.extend([
                    "--anonymous-diarization-pipeline", args.anonymous_diarization_pipeline,
                    "--anonymous-device", args.anonymous_device,
                ])
            for model in args.identity_model:
                command.extend(["--identity-model", model])
            if args.overwrite:
                command.append("--overwrite")
            completed = subprocess.run(command, check=False)
            if completed.returncode in {0, 2} and (package / "handoff_manifest.json").is_file():
                if not args.out_dir:
                    (script_root / "02_script_work").mkdir(parents=True, exist_ok=True)
                    (script_root / "03_final_delivery").mkdir(parents=True, exist_ok=True)
                discovery_dir = package / "06_handoff" / "material_discovery"
                discovery_dir.mkdir(parents=True, exist_ok=True)
                (discovery_dir / discovery_json.name).write_text(discovery_json.read_text(encoding="utf-8"), encoding="utf-8")
                write_manifest(discovery_dir / "episode_manifest.tsv", result["episode_rows"])
            return completed.returncode
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
