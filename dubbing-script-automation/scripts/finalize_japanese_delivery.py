#!/usr/bin/env python3
"""Build the three Japanese client-delivery Excel workbooks.

The QC-approved episode TSVs are authoritative for role, timecode, and
dialogue. Approved Word scripts are an independently rendered checkpoint and
must match those TSV values line for line before export.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import unicodedata
import uuid
from collections import Counter
from pathlib import Path
from typing import Any, Callable

from dubbing_tool import qc_snapshot_issues

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple


TSV_FIELDS = [
    "episode", "source_index", "start", "end", "role", "confidence",
    "status", "frame", "text", "provider", "model",
]
LEGACY_ROLE_MAP = {"旁白音声": "ナレーション"}
GENERIC_ROLE_PREFIXES = ("男性音声", "女性音声", "子供音声", "老人男性音声", "老人女性音声", "ナレーション", "不明音声")
CHARACTER_HEADERS = [
    "名前", "キャラクター画像", "役割/身分", "性別", "年齢",
    "キャラクター説明", "声のトーン/声質",
]


class DeliveryError(RuntimeError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def file_snapshot(path: Path) -> dict[str, Any]:
    return {
        "path": str(path.resolve()),
        "sha256": sha256_file(path),
        "bytes": path.stat().st_size,
    }


def verify_qc_snapshot(directory: Path) -> dict[str, Any]:
    """Ensure finalization starts from the exact TSV that passed episode QC."""
    tsv_path = directory / "final_roles.tsv"
    report_path = directory / "qc_report.json"
    if not tsv_path.is_file():
        raise DeliveryError(f"Missing final role TSV: {tsv_path}")
    if not report_path.is_file():
        raise DeliveryError(f"Missing QC report: {report_path}")
    try:
        report = json.loads(report_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise DeliveryError(f"Cannot read QC report: {report_path}") from exc
    if report.get("production_ready") is not True:
        raise DeliveryError(f"Episode {directory.name} is not QC-approved according to {report_path}")
    snapshot = report.get("artifact_snapshot")
    tsv_snapshot = snapshot.get("tsv") if isinstance(snapshot, dict) else None
    expected = str(tsv_snapshot.get("sha256", "")) if isinstance(tsv_snapshot, dict) else ""
    if not expected:
        raise DeliveryError(f"QC snapshot is missing for episode {directory.name}: {report_path}")
    actual = sha256_file(tsv_path)
    if actual.casefold() != expected.casefold():
        raise DeliveryError(f"QC snapshot is stale for episode {directory.name}: final_roles.tsv changed after QC")
    source_issues = qc_snapshot_issues(report, tsv_path, require_source=True)
    if source_issues:
        raise DeliveryError(f"Episode {directory.name} source QC is stale or incomplete: {source_issues}")
    return {
        "qc_report": file_snapshot(report_path),
        "final_roles": file_snapshot(tsv_path),
    }


def verify_user_confirmation_locks(
    directory: Path, rows: list[dict[str, str]],
) -> dict[str, Any] | None:
    """Fail if a human-confirmed role was changed or lost downstream."""
    checkpoint_candidates = [
        directory / f"episode{directory.name}_checkpoint.json",
        directory / "episode_checkpoint.json",
    ]
    checkpoint_path = next((path for path in checkpoint_candidates if path.is_file()), None)
    if checkpoint_path is None:
        return None
    try:
        checkpoint = json.loads(checkpoint_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise DeliveryError(f"Cannot read episode checkpoint: {checkpoint_path}") from exc
    if checkpoint.get("user_mapping_confirmed") is not True:
        return None

    manifest_path = directory / "user_confirmation_manifest.json"
    if not manifest_path.is_file():
        raise DeliveryError(
            f"Human-confirmed episode {directory.name} is missing user_confirmation_manifest.json"
        )
    try:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise DeliveryError(f"Cannot read user confirmation manifest: {manifest_path}") from exc
    source = manifest.get("source_workbook") or {}
    source_path = Path(str(source.get("path", "")))
    expected_hash = str(source.get("sha256", ""))
    if not source_path.is_file():
        raise DeliveryError(
            f"User-approved workbook recorded for episode {directory.name} is missing: {source_path}"
        )
    if not expected_hash or sha256_file(source_path).casefold() != expected_hash.casefold():
        raise DeliveryError(
            f"User-approved workbook hash changed for episode {directory.name}: {source_path}"
        )

    by_source_index: dict[str, dict[str, str]] = {}
    for position, row in enumerate(rows, start=1):
        source_index = str(row.get("source_index", "")).strip() or str(position)
        if source_index in by_source_index:
            raise DeliveryError(
                f"Duplicate source_index {source_index!r} in final_roles.tsv for episode {directory.name}"
            )
        by_source_index[source_index] = row
    approved_rows = manifest.get("approved_rows")
    if not isinstance(approved_rows, list) or not approved_rows:
        raise DeliveryError(
            f"User confirmation manifest has no approved rows for episode {directory.name}"
        )
    if len(approved_rows) != len(rows):
        raise DeliveryError(
            f"User lock row-count mismatch for episode {directory.name}: "
            f"manifest={len(approved_rows)}, final_roles={len(rows)}"
        )
    for approved in approved_rows:
        source_index = str(approved.get("source_index", "")).strip()
        row = by_source_index.get(source_index)
        if row is None:
            raise DeliveryError(
                f"User-locked source row {source_index!r} is missing in episode {directory.name}"
            )
        approved_role = str(approved.get("role", ""))
        if str(row.get("user_locked", "")).strip().casefold() != "true":
            raise DeliveryError(
                f"User lock flag is missing in episode {directory.name}, source row {source_index}"
            )
        if str(row.get("user_locked_role", "")) != approved_role:
            raise DeliveryError(
                f"Stored user lock differs from manifest in episode {directory.name}, "
                f"source row {source_index}"
            )
        if str(row.get("role", "")) != approved_role:
            raise DeliveryError(
                f"Human-approved role was overwritten in episode {directory.name}, "
                f"source row {source_index}: approved={approved_role!r}, "
                f"final={row.get('role', '')!r}"
            )
    return {
        "checkpoint": file_snapshot(checkpoint_path),
        "manifest": file_snapshot(manifest_path),
        "approved_workbook": file_snapshot(source_path),
        "locked_rows": len(approved_rows),
    }


def normalized(value: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", value).split())


def first_line(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n").split("\n", 1)[0].strip()


def normalized_line_endings(value: str) -> str:
    """Normalize container line endings without changing dialogue content."""
    return value.replace("\r\n", "\n").replace("\r", "\n")


def has_japanese(value: str) -> bool:
    return any("\u3040" <= char <= "\u30ff" or "\u3400" <= char <= "\u9fff" for char in value)


def role_is_generic(role: str) -> bool:
    return normalized(role).startswith(GENERIC_ROLE_PREFIXES)


def episode_dirs(root: Path) -> list[Path]:
    directories = [path for path in root.iterdir() if path.is_dir() and re.fullmatch(r"\d{4}", path.name)]
    if not directories:
        raise DeliveryError(f"No numbered episode folders found under: {root}")
    return sorted(directories, key=lambda path: path.name)


def read_tsv(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise DeliveryError(f"Missing final role TSV: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter="\t")
        if set(TSV_FIELDS) - set(reader.fieldnames or []):
            raise DeliveryError(f"Invalid role TSV header: {path}")
        return list(reader)


def write_tsv(path: Path, rows: list[dict[str, str]]) -> None:
    fields = list(TSV_FIELDS)
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def collect_named_roles(episodes_root: Path) -> set[str]:
    roles: set[str] = set()
    for directory in episode_dirs(episodes_root):
        for row in read_tsv(directory / "final_roles.tsv"):
            role = row.get("role", "").strip()
            if role and not role_is_generic(role):
                roles.add(role)
    return roles


def extract_name_map(path: Path, named_roles: set[str]) -> dict[str, str]:
    """Find English-name/Japanese-name pairs in a supplied workbook.

    Supports direct two-column tables and title/subtitle sheets where an
    English character name begins one cell and its Japanese translation is in
    a neighboring cell.
    """
    if not path.is_file():
        raise DeliveryError(f"Japanese name workbook not found: {path}")
    role_lookup = {normalized(role).casefold(): role for role in named_roles}
    workbook = load_workbook(path, read_only=True, data_only=True)
    mapping: dict[str, str] = {}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [str(cell).strip() if cell is not None else "" for cell in row]
            # An explicit two-column mapping may preserve an approved official Roman name.
            if len(values) == 2:
                source = role_lookup.get(normalized(values[0]).casefold())
                if source and values[1]:
                    mapping[source] = values[1]
                    continue
            for index, value in enumerate(values):
                source = role_lookup.get(normalized(first_line(value)).casefold())
                if not source:
                    continue
                candidates = values[index + 1 :] + values[:index]
                target = next((first_line(candidate) for candidate in candidates if has_japanese(first_line(candidate))), "")
                if target:
                    mapping[source] = target
    return mapping


def map_role(role: str, name_map: dict[str, str]) -> str:
    role = LEGACY_ROLE_MAP.get(role.strip(), role.strip())
    return name_map.get(role, role)


def parse_word_script(path: Path, episode: str) -> tuple[str, list[tuple[str, str, str]]]:
    document = Document(path)
    if not document.paragraphs:
        raise DeliveryError(f"Word script is empty: {path}")
    title = document.paragraphs[0].text.strip().replace("集", "話") or f"第{episode}話"
    rows: list[tuple[str, str, str]] = []
    for paragraph in document.paragraphs[1:]:
        parts = paragraph.text.split("\t")
        if len(parts) != 3:
            continue
        role, timecode, text = parts
        role = "" if role.strip() in {"", "　"} else role.strip()
        rows.append((role, timecode, text))
    return title, rows


def reconcile_word_and_tsv(
    episode: str,
    lines: list[tuple[str, str, str]],
    rows: list[dict[str, str]],
    name_map: dict[str, str],
) -> list[tuple[str, str, str]]:
    """Validate the Word checkpoint and return TSV-authoritative output rows."""
    if len(lines) != len(rows):
        raise DeliveryError(
            f"Word/TSV line-count mismatch in episode {episode}: "
            f"Word={len(lines)}, TSV={len(rows)}"
        )
    reconciled: list[tuple[str, str, str]] = []
    prior_role = ""
    for position, (role, timecode, text) in enumerate(lines):
        source_role = map_role(role, name_map) if role else ""
        if source_role:
            prior_role = source_role
        expected_role = rows[position]["role"]
        if prior_role != expected_role:
            raise DeliveryError(
                f"Word/TSV role mismatch in episode {episode}, line {position + 1}: "
                f"Word={prior_role!r}, TSV={expected_role!r}"
            )
        expected_timecode = rows[position]["start"]
        if timecode != expected_timecode:
            raise DeliveryError(
                f"Word/TSV timecode mismatch in episode {episode}, line {position + 1}: "
                f"Word={timecode!r}, TSV={expected_timecode!r}"
            )
        expected_text = normalized_line_endings(rows[position]["text"])
        word_text = normalized_line_endings(text)
        if word_text != expected_text:
            raise DeliveryError(
                f"Word/TSV dialogue mismatch in episode {episode}, line {position + 1}: "
                f"Word={word_text!r}, TSV={expected_text!r}"
            )
        reconciled.append((source_role, expected_timecode, expected_text))
    return reconciled


def style_sheet(sheet: Any, header_row: int, end_row: int, end_column: int) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    thin = Side(style="thin", color="D9E2F3")
    for cell in sheet[header_row]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=header_row, max_row=end_row, min_col=1, max_col=end_column):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for row in range(header_row + 1, end_row + 1):
        if sheet.cell(row, 1).value and sheet.cell(row, 2).value is None and sheet.cell(row, 3).value is None:
            for column in range(1, end_column + 1):
                sheet.cell(row, column).fill = section_fill
            sheet.cell(row, 1).font = Font(bold=True, color="1F4E78", size=12)


def build_script_workbook(records: list[tuple[str, list[tuple[str, str, str]]]], output: Path) -> list[list[Any]]:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "総合脚本"
    sheet.sheet_view.showGridLines = False
    sheet.append(["名前", "タイムコード", "台詞"])
    expected: list[list[Any]] = [["名前", "タイムコード", "台詞"]]
    for title, lines in records:
        row = sheet.max_row + 1
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sheet.cell(row, 1, title)
        expected.append([title, None, None])
        for role, timecode, text in lines:
            sheet.append([role, timecode, text])
            expected.append([role or None, timecode, text])
    style_sheet(sheet, 1, sheet.max_row, 3)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 70
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 23 if sheet.cell(row, 2).value is None else 30
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return expected


def build_cast_workbook(episodes: dict[str, list[dict[str, str]]], output: Path) -> None:
    all_roles: list[str] = []
    for rows in episodes.values():
        for row in rows:
            role = row["role"].strip()
            if role and role not in all_roles:
                all_roles.append(role)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "香盤表"
    sheet.sheet_view.showGridLines = False
    episode_names = list(episodes)
    episode_headers = [f"第{episode}話" for episode in episode_names]
    sheet.append(["役名", *episode_headers, "総出演話数"])
    for role in all_roles:
        appearances = ["○" if any(row["role"].strip() == role for row in episodes[episode]) else "" for episode in episode_names]
        sheet.append([role, *appearances, sum(value == "○" for value in appearances)])
    style_sheet(sheet, 1, sheet.max_row, sheet.max_column)
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions[get_column_letter(sheet.max_column)].width = 14
    for column in range(2, sheet.max_column):
        sheet.column_dimensions[get_column_letter(column)].width = 7
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def image_anchor_cell(image: Any) -> tuple[int, int] | None:
    anchor = image.anchor
    marker = getattr(anchor, "_from", None)
    if marker is not None:
        return int(marker.row) + 1, int(marker.col) + 1
    if isinstance(anchor, str):
        return coordinate_to_tuple(anchor)
    return None


def validate_character_workbook(
    path: Path, require_delivery_sheet_name: bool = False,
) -> dict[str, Any]:
    if not path.is_file():
        raise DeliveryError(f"Character-settings workbook not found: {path}")
    workbook = load_workbook(path, data_only=False, read_only=False)
    if len(workbook.worksheets) != 1:
        raise DeliveryError(
            "Character-settings workbook must contain exactly one worksheet"
        )
    sheet = workbook.active
    if require_delivery_sheet_name and sheet.title != "登場人物設定":
        raise DeliveryError(
            f"Character-settings worksheet must be named 登場人物設定: {sheet.title!r}"
        )
    headers = [sheet.cell(1, column).value for column in range(1, 8)]
    if headers != CHARACTER_HEADERS:
        raise DeliveryError(
            "Character-settings headers must be exactly: "
            + " | ".join(CHARACTER_HEADERS)
        )
    image_rows = {
        position[0]
        for image in sheet._images
        for position in [image_anchor_cell(image)]
        if position is not None and position[1] == 2
    }
    names: list[str] = []
    if any(cell.value not in (None, "") for row in sheet.iter_rows(min_col=8) for cell in row):
        raise DeliveryError("Character-settings workbook contains extra columns outside the seven-column contract")
    rows: list[int] = []
    for row in range(2, sheet.max_row + 1):
        values = [sheet.cell(row, column).value for column in range(1, 8)]
        if not any(value not in (None, "") for value in values):
            continue
        rows.append(row)
        names.append(str(sheet.cell(row, 1).value or "").strip())
        required_columns = (1, 3, 4, 5, 6, 7)
        missing = [
            CHARACTER_HEADERS[column - 1]
            for column in required_columns
            if sheet.cell(row, column).value in (None, "")
        ]
        if missing:
            raise DeliveryError(
                f"Character-settings row {row} is missing: {', '.join(missing)}"
            )
        if row not in image_rows:
            raise DeliveryError(
                f"Character-settings row {row} has no portrait anchored in column B"
            )
        for column in (3, 6, 7):
            value = str(sheet.cell(row, column).value or "")
            if not has_japanese(value):
                raise DeliveryError(
                    f"Character-settings row {row} column {CHARACTER_HEADERS[column - 1]} must contain Japanese text"
                )
        gender = normalized(str(sheet.cell(row, 4).value or ""))
        if gender not in {"男", "女", "その他", "不明"}:
            raise DeliveryError(
                f"Character-settings row {row} has a non-Japanese 性別 value: {gender!r}"
            )
    if not rows:
        raise DeliveryError("Character-settings workbook contains no character rows")
    if len(set(names)) != len(names):
        raise DeliveryError("Character-settings workbook contains duplicate names")
    return {
        "worksheet": sheet.title,
        "characters": len(rows),
        "names": names,
        "portrait_rows": len(image_rows.intersection(rows)),
    }


def copy_character_workbook(source: Path, output: Path) -> dict[str, Any]:
    validate_character_workbook(source)
    shutil.copy2(source, output)
    workbook = load_workbook(output, data_only=False, read_only=False)
    workbook.active.title = "登場人物設定"
    workbook.save(output)
    return validate_character_workbook(output, require_delivery_sheet_name=True)


def validate_formal_output_target(
    output_dir: Path, filenames: list[str], overwrite: bool,
) -> None:
    output_dir.parent.mkdir(parents=True, exist_ok=True)
    if not output_dir.exists():
        return
    if not output_dir.is_dir():
        raise DeliveryError(f"Formal delivery target is not a directory: {output_dir}")
    allowed = set(filenames)
    extras = [
        path for path in output_dir.iterdir()
        if not path.name.startswith(".") and path.name not in allowed
    ]
    if extras:
        raise DeliveryError(
            "Formal delivery directory contains files outside the three-file contract: "
            + ", ".join(sorted(path.name for path in extras))
        )
    for filename in filenames:
        target = output_dir / filename
        if not target.exists():
            continue
        if not overwrite:
            raise DeliveryError(
                f"Formal delivery output already exists: {target}. Use --overwrite."
            )
        if not target.is_file():
            raise DeliveryError(f"Formal delivery target is not a file: {target}")


def commit_staged_output(
    staging_dir: Path,
    output_dir: Path,
    post_commit: Callable[[], None] | None = None,
) -> None:
    """Replace delivery atomically and restore the old directory on any late failure."""
    backup_dir: Path | None = None
    new_output_installed = False
    try:
        if output_dir.exists():
            backup_dir = output_dir.parent / f".{output_dir.name}.backup-{uuid.uuid4().hex}"
            os.replace(output_dir, backup_dir)
        os.replace(staging_dir, output_dir)
        new_output_installed = True
        if post_commit is not None:
            post_commit()
    except Exception:
        if new_output_installed and output_dir.exists():
            if output_dir.is_dir():
                shutil.rmtree(output_dir)
            else:
                output_dir.unlink()
        if backup_dir is not None and backup_dir.exists() and not output_dir.exists():
            os.replace(backup_dir, output_dir)
        raise
    else:
        if backup_dir is not None and backup_dir.exists():
            shutil.rmtree(backup_dir)


def verify_workbook(path: Path, expected: list[list[Any]]) -> int:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["総合脚本"]
    actual = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    if actual != expected:
        for index, (wanted, found) in enumerate(zip(expected, actual), start=1):
            if wanted != found:
                raise DeliveryError(f"Japanese script workbook differs from TSV-authoritative data at row {index}: expected={wanted!r}, actual={found!r}")
        raise DeliveryError("Japanese script workbook row count differs from TSV-authoritative data")
    return len(actual) - 1


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Create exactly three Japanese production Excel workbooks from approved internal work artifacts."
    )
    parser.add_argument("--episodes-root", required=True, help="Folder containing 0000, 0001, ... episode folders")
    parser.add_argument(
        "--work-dir", "--delivery-dir", dest="work_dir", required=True,
        help="Clean internal assemble-delivery work folder containing approved_scripts",
    )
    parser.add_argument(
        "--out-dir", required=True,
        help="Formal output folder; it must contain only the three Japanese Excel deliverables",
    )
    parser.add_argument("--name-map-xlsx", required=True, help="Workbook containing source-name/Japanese-name pairs")
    parser.add_argument(
        "--character-xlsx", required=True,
        help="Finalized client-facing character workbook with portraits and the required seven Japanese columns",
    )
    parser.add_argument("--out-script", default="日本語_総合脚本.xlsx")
    parser.add_argument(
        "--out-kouban", "--out-cast", dest="out_kouban",
        default="日本語_香盤表.xlsx",
    )
    parser.add_argument("--out-characters", default="日本語_登場人物設定表.xlsx")
    parser.add_argument(
        "--report",
        help="Internal QC report path; defaults under --work-dir and may not be inside --out-dir",
    )
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    episodes_root = Path(args.episodes_root)
    work_dir = Path(args.work_dir)
    output_dir = Path(args.out_dir)
    character_source = Path(args.character_xlsx)
    output_names = [args.out_script, args.out_kouban, args.out_characters]
    if len(set(output_names)) != 3:
        raise DeliveryError("The three formal output filenames must be distinct")
    for name in output_names:
        if Path(name).name != name or Path(name).suffix.casefold() != ".xlsx":
            raise DeliveryError(
                f"Formal output names must be plain .xlsx filenames: {name!r}"
            )
        if not has_japanese(Path(name).stem):
            raise DeliveryError(
                f"Formal output filenames must include Japanese text: {name!r}"
            )
    if work_dir.resolve() == output_dir.resolve() or output_dir.resolve() in work_dir.resolve().parents:
        raise DeliveryError("Formal output and internal work roots must be separate")
    protected = {character_source.resolve(), Path(args.name_map_xlsx).resolve()}
    if any((output_dir / name).resolve() in protected for name in output_names):
        raise DeliveryError("A formal output would overwrite an input workbook")
    character_summary = validate_character_workbook(character_source)
    named_roles = collect_named_roles(episodes_root)
    name_map = extract_name_map(Path(args.name_map_xlsx), named_roles)
    translated_roles = {name_map.get(role, role) for role in named_roles}
    unmapped = sorted(role for role in named_roles if not has_japanese(role) and role not in name_map)
    if unmapped:
        raise DeliveryError(f"Japanese name mapping is incomplete for named roles: {unmapped}")

    missing_characters = sorted(translated_roles - set(character_summary["names"]))
    if missing_characters:
        raise DeliveryError(f"Character-settings workbook is missing script characters: {missing_characters}")
    summary_path = work_dir / "overall_qc_summary.json"
    if not summary_path.is_file():
        raise DeliveryError("Delivery work is missing the complete-series assembly summary")
    summary = json.loads(summary_path.read_text(encoding="utf-8"))
    current_episodes = {path.name for path in episode_dirs(episodes_root)}
    if summary.get("complete") is not True or set(summary.get("expected_episode_ids") or []) != current_episodes:
        raise DeliveryError("Rebuild delivery work with --expected-episodes and --require-complete; episode coverage is incomplete or stale")
    script_dir = work_dir / "approved_scripts"
    if not script_dir.is_dir():
        raise DeliveryError(
            f"Approved Word scripts are missing: {script_dir}. "
            "Rebuild delivery_work without --include-review."
        )
    documents = {path.name[:4]: path for path in script_dir.glob("[0-9][0-9][0-9][0-9]_*.docx")}
    if set(path.name for path in episode_dirs(episodes_root)) - set(documents):
        missing = sorted(set(path.name for path in episode_dirs(episodes_root)) - set(documents))
        raise DeliveryError(f"Word scripts are missing for episodes: {missing}")

    source_qc_snapshots: dict[str, dict[str, Any]] = {}
    user_confirmation_snapshots: dict[str, dict[str, Any]] = {}
    for directory in episode_dirs(episodes_root):
        rows = read_tsv(directory / "final_roles.tsv")
        source_qc_snapshots[directory.name] = verify_qc_snapshot(directory)
        confirmation_snapshot = verify_user_confirmation_locks(directory, rows)
        if confirmation_snapshot is not None:
            user_confirmation_snapshots[directory.name] = confirmation_snapshot
    localized_roles_dir = work_dir / "formal_delivery_work" / "role_tsv"
    episodes: dict[str, list[dict[str, str]]] = {}
    if not args.dry_run:
        for directory in episode_dirs(episodes_root):
            rows = read_tsv(directory / "final_roles.tsv")
            for row in rows:
                row["role"] = map_role(row["role"], name_map)
            write_tsv(localized_roles_dir / f"{directory.name}_final_roles.tsv", rows)
            episodes[directory.name] = rows
    else:
        for directory in episode_dirs(episodes_root):
            rows = read_tsv(directory / "final_roles.tsv")
            for row in rows:
                row["role"] = map_role(row["role"], name_map)
            episodes[directory.name] = rows

    records: list[tuple[str, list[tuple[str, str, str]]]] = []
    word_count = 0
    for episode, rows in episodes.items():
        title, lines = parse_word_script(documents[episode], episode)
        normalized_lines = reconcile_word_and_tsv(episode, lines, rows, name_map)
        records.append((f"第{episode}話", normalized_lines))
        word_count += len(normalized_lines)

    report = {
        "episodes": len(records),
        "word_rows": word_count,
        "name_map": name_map,
        "unmapped_named_roles": unmapped,
        "character_workbook": character_summary,
        "dry_run": args.dry_run,
        "source_qc_snapshots": source_qc_snapshots,
        "user_confirmation_snapshots": user_confirmation_snapshots,
    }
    if not args.dry_run:
        report_path = (
            Path(args.report)
            if args.report
            else work_dir / "formal_delivery_qc" / "納品チェック.json"
        )
        try:
            report_path.resolve().relative_to(output_dir.resolve())
        except ValueError:
            pass
        else:
            raise DeliveryError(
                "Internal QC report may not be written inside the formal delivery directory"
            )
        validate_formal_output_target(output_dir, output_names, args.overwrite)
        staging_dir = Path(tempfile.mkdtemp(
            prefix=f".{output_dir.name}.staging-", dir=output_dir.parent,
        ))
        try:
            staged_script = staging_dir / args.out_script
            staged_kouban = staging_dir / args.out_kouban
            staged_characters = staging_dir / args.out_characters
            expected = build_script_workbook(records, staged_script)
            report["excel_rows"] = verify_workbook(staged_script, expected) - len(records)
            build_cast_workbook(episodes, staged_kouban)
            report["character_workbook"] = copy_character_workbook(
                character_source, staged_characters,
            )
            staged_entries = sorted(
                path.name for path in staging_dir.iterdir()
                if not path.name.startswith(".")
            )
            if staged_entries != sorted(output_names):
                raise DeliveryError(
                    "Staged formal delivery does not contain exactly three expected files: "
                    f"{staged_entries}"
                )
            def finish_commit() -> None:
                script_path = output_dir / args.out_script
                kouban_path = output_dir / args.out_kouban
                character_path = output_dir / args.out_characters
                report["script"] = str(script_path)
                report["koubanhyo"] = str(kouban_path)
                report["characters"] = str(character_path)
                report["word_tsv_diff_count"] = 0
                report["tsv_excel_diff_count"] = 0
                artifact_snapshot: dict[str, Any] = {}
                for directory in episode_dirs(episodes_root):
                    role_tsv = localized_roles_dir / f"{directory.name}_final_roles.tsv"
                    artifact_snapshot[f"role_tsv/{role_tsv.name}"] = file_snapshot(role_tsv)
                    document = documents[directory.name]
                    artifact_snapshot[f"word/{document.name}"] = file_snapshot(document)
                artifact_snapshot["総合脚本"] = file_snapshot(script_path)
                artifact_snapshot["香盤表"] = file_snapshot(kouban_path)
                artifact_snapshot["登場人物設定表"] = file_snapshot(character_path)
                report["artifact_snapshot"] = {
                    "schema_version": 3,
                    "files": artifact_snapshot,
                }
                final_entries = sorted(
                    path.name for path in output_dir.iterdir()
                    if not path.name.startswith(".")
                )
                if final_entries != sorted(output_names):
                    raise DeliveryError(
                        "Formal delivery directory does not contain exactly three expected files: "
                        f"{final_entries}"
                    )
                report["formal_delivery_files"] = final_entries
                report["internal_qc_report"] = str(report_path)
                report_path.parent.mkdir(parents=True, exist_ok=True)
                descriptor, temporary_report_name = tempfile.mkstemp(
                    prefix=f".{report_path.name}.", suffix=".tmp", dir=report_path.parent,
                )
                os.close(descriptor)
                temporary_report = Path(temporary_report_name)
                try:
                    temporary_report.write_text(
                        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8",
                    )
                    os.replace(temporary_report, report_path)
                finally:
                    if temporary_report.exists():
                        temporary_report.unlink()

            commit_staged_output(staging_dir, output_dir, post_commit=finish_commit)
        finally:
            if staging_dir.exists():
                shutil.rmtree(staging_dir)
    print(json.dumps(report, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except DeliveryError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(2)
